"""Document ingestion, retrieval, and DeepSeek-powered answering for RAG Web."""

from __future__ import annotations

import os
import re
import uuid
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import chromadb
from chromadb.api.models.Collection import Collection
from dotenv import load_dotenv
from openai import OpenAI
from pypdf import PdfReader
from docx import Document as DocxDocument

load_dotenv()

# ===== 使用智谱 Embedding API（不再本地加载模型，省内存） =====
ZHIPU_BASE_URL = "https://open.bigmodel.cn/api/paas/v4/"
EMBEDDING_MODEL = os.getenv("EMBEDDING_MODEL", "embedding-3")
EMBEDDING_DIMENSIONS = int(os.getenv("EMBEDDING_DIMENSIONS", "1024"))
EMBEDDING_BATCH_SIZE = 25

DEEPSEEK_BASE_URL = "https://api.deepseek.com"
DEFAULT_TOP_K = max(1, int(os.getenv("RAG_TOP_K", "8")))


@dataclass
class Document:
    """A retrieved text chunk and its source metadata."""

    content: str
    source: str
    chunk_index: int = 0


def read_file(file_path: str) -> str:
    """读取文档，支持 txt、md、pdf、docx，docx 遇到图片错误时跳过"""
    path = Path(file_path)
    ext = path.suffix.lower()
    if ext == ".txt":
        for encoding in ["utf-8", "gbk", "gb2312", "latin-1"]:
            try:
                with path.open("r", encoding=encoding) as file:
                    return file.read()
            except UnicodeDecodeError:
                continue
        raise ValueError(f"无法解码文件: {file_path}")
    if ext == ".md":
        return path.read_text(encoding="utf-8")
    if ext == ".pdf":
        return "\n\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if ext == ".docx":
        try:
            doc = DocxDocument(str(path))
            return "\n\n".join(p.text for p in doc.paragraphs)
        except Exception as e:
            print(f"读取 docx 文件失败 ({file_path}): {e}，尝试只读取文本内容")
            # 用 zipfile 读取 docx 中的 document.xml 提取文本
            import zipfile
            import xml.etree.ElementTree as ET
            try:
                with zipfile.ZipFile(path, 'r') as zf:
                    with zf.open('word/document.xml') as xml_file:
                        tree = ET.parse(xml_file)
                        root = tree.getroot()
                        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                        texts = []
                        for t in root.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'):
                            if t.text:
                                texts.append(t.text)
                        return "\n\n".join(texts)
            except Exception as e2:
                print(f"docx 文本提取也失败: {e2}")
                return ""
    raise ValueError("仅支持 txt、md、pdf 和 docx 格式的文档")


class RAGEngine:
    """Stores paragraph chunks in ChromaDB and answers questions from retrieved context."""

    def __init__(
        self,
        persist_directory: str | Path | None = None,
        collection_name: str = "rag_documents_zhipu_v1",
    ) -> None:
        database_path = Path(
            persist_directory
            or os.getenv("CHROMA_DB_PATH", Path(__file__).parent / "chroma_data")
        )
        database_path.mkdir(parents=True, exist_ok=True)
        self._client = chromadb.PersistentClient(path=str(database_path))
        self._collection: Collection = self._client.get_or_create_collection(
            name=collection_name,
            metadata={"hnsw:space": "cosine"},
        )
        self._embed_client: OpenAI | None = None

    @property
    def embed_client(self) -> OpenAI:
        """Lazily create the Zhipu embedding client (OpenAI-compatible endpoint)."""
        if self._embed_client is None:
            api_key = os.getenv("ZHIPU_API_KEY")
            if not api_key:
                raise ValueError("请设置 ZHIPU_API_KEY 环境变量")
            self._embed_client = OpenAI(api_key=api_key, base_url=ZHIPU_BASE_URL)
        return self._embed_client

    def _embed_texts(self, texts: list[str]) -> list[list[float]]:
        """调用智谱 Embedding API，分批把文本转换为向量，避免单次请求过大。"""
        if not texts:
            return []
        embeddings: list[list[float]] = []
        for start in range(0, len(texts), EMBEDDING_BATCH_SIZE):
            batch = texts[start:start + EMBEDDING_BATCH_SIZE]
            response = self.embed_client.embeddings.create(
                model=EMBEDDING_MODEL,
                input=batch,
                dimensions=EMBEDDING_DIMENSIONS,
            )
            embeddings.extend(item.embedding for item in response.data)
        return embeddings

    def load_document(self, file_path: str | Path) -> list[str]:
        """Read a TXT, Markdown, PDF, or DOCX file and return paragraph chunks."""
        path = Path(file_path)
        if not path.is_file():
            raise FileNotFoundError(f"找不到文档：{path}")

        return self.chunk_text(read_file(str(path)))

    @staticmethod
    def chunk_text(text: str) -> list[str]:
        """Split text into non-empty chunks based on blank-line-separated paragraphs."""
        return [
            re.sub(r"\s+", " ", paragraph).strip()
            for paragraph in re.split(r"\n\s*\n+", text)
            if paragraph.strip()
        ]

    def ingest_file(self, file_path: str | Path, source: str | None = None) -> int:
        """Load a supported document, split it by paragraph, and index every chunk."""
        path = Path(file_path)
        chunks = self.load_document(path)
        self._store_chunks(chunks, source=source or path.name)
        return len(chunks)

    def add_document(self, content: str, source: str = "local") -> int:
        """Index plain text content; retained for API clients that submit text directly."""
        chunks = self.chunk_text(content)
        self._store_chunks(chunks, source=source)
        return len(chunks)

    def _store_chunks(self, chunks: list[str], source: str) -> None:
        if not chunks:
            return
        embeddings = self._embed_texts(chunks)
        ids = [str(uuid.uuid4()) for _ in chunks]
        metadatas = [{"source": source, "chunk_index": index} for index in range(len(chunks))]
        self._collection.add(
            ids=ids,
            documents=chunks,
            embeddings=embeddings,
            metadatas=metadatas,
        )

    def list_documents(self) -> list[dict[str, int | str]]:
        """Return the indexed files and their number of stored chunks."""
        metadata = self._collection.get(include=["metadatas"]).get("metadatas", [])
        counts = Counter(str(item.get("source", "unknown")) for item in metadata if item)
        return [
            {"source": source, "chunks": chunk_count}
            for source, chunk_count in sorted(counts.items())
        ]

    def get_document_chunks(self, source: str) -> list[str]:
        """获取某个文件的所有分块内容（用于预览）"""
        results = self._collection.get(
            where={"source": source},
            include=["documents"]
        )
        return results.get("documents", [])

    def get_document_chunk(self, source: str, chunk_index: int) -> Document | None:
        """Fetch one original chunk so a citation can be opened in the UI."""
        results = self._collection.get(
            where={"source": source},
            include=["documents", "metadatas"],
        )
        documents = results.get("documents", [])
        metadatas = results.get("metadatas", [])
        for content, metadata in zip(documents, metadatas):
            if content and metadata and int(metadata.get("chunk_index", -1)) == chunk_index:
                return Document(content=content, source=source, chunk_index=chunk_index)
        return None

    def retrieve(self, query: str, top_k: int | None = None) -> list[Document]:
        """Return several semantically related chunks, eight by default."""
        if not query.strip() or self._collection.count() == 0:
            return []
        limit = top_k or DEFAULT_TOP_K
        query_embedding = self._embed_texts([query])
        results = self._collection.query(
            query_embeddings=query_embedding,
            n_results=min(limit, self._collection.count()),
            include=["documents", "metadatas"],
        )
        documents = results.get("documents", [[]])[0]
        metadatas = results.get("metadatas", [[]])[0]
        return [
            Document(
                content=content,
                source=str(metadata.get("source", "unknown")),
                chunk_index=int(metadata.get("chunk_index", 0)),
            )
            for content, metadata in zip(documents, metadatas)
            if content and metadata
        ]

    def answer(self, query: str, top_k: int | None = None) -> tuple[str, list[Document]]:
        """Retrieve context and ask DeepSeek to produce a grounded Chinese response."""
        documents = self.retrieve(query, top_k=top_k)
        if not documents:
            return "没有找到相关内容。请先添加文档，或换一种问法。", []

        api_key = os.getenv("DEEPSEEK_API_KEY")
        if not api_key:
            raise ValueError("请设置 DEEPSEEK_API_KEY 环境变量")

        context = "\n\n".join(document.content for document in documents)
        client = OpenAI(api_key=api_key, base_url=DEEPSEEK_BASE_URL)
        response = client.chat.completions.create(
            model="deepseek-chat",
            messages=[
                {
                    "role": "system",
                    "content": (
                        "你是严谨的中文知识库助手。仅依据给出的上下文回答，不要编造。"
                        "请综合所有相关片段形成完整答案，而不是只复述其中一段；"
                        "若片段信息相互补充，请明确整合；若上下文不足或有矛盾，请明确说明。"
                        "使用清晰的中文作答，不要使用任何加粗或特殊格式。"
                    ),
                },
                {
                    "role": "user",
                    "content": f"上下文：\n{context}\n\n问题：{query}",
                },
            ],
            temperature=0.2,
        )
        return response.choices[0].message.content or "未能生成回答。", documents